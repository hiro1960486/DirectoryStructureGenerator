import csv
import os
import tempfile
import unittest
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

from dsg_app.file_inspector import FileDetail
from dsg_app.organizer import (
    RESULT_COLUMNS, RESULT_LOG_NAME, CopyCancelled, build_copy_plans,
    classify_copy_error, execute_copy_plans, excel_hyperlink, windows_file_uri,
)


def detail(path: Path, relative: str) -> FileDetail:
    return FileDetail(
        name=path.name, relative_path=relative, full_path=str(path),
        category="general", extension=path.suffix, file_size=path.stat().st_size,
        modified="2026-09-23 10:00:00", created="2026-09-23 09:00:00",
        permissions="-rw-r--r--", readonly=False,
    )


class WindowsLinkTests(unittest.TestCase):
    def test_drive_path_is_percent_encoded_uri(self):
        self.assertEqual(
            windows_file_uri(r"c:\写真\2026 09"),
            "file:///C:/%E5%86%99%E7%9C%9F/2026%2009",
        )

    def test_unc_path_is_percent_encoded_uri(self):
        self.assertEqual(
            windows_file_uri(r"\\server\share\写真"),
            "file://server/share/%E5%86%99%E7%9C%9F",
        )

    def test_excel_formula_prefers_native_windows_path(self):
        self.assertEqual(
            excel_hyperlink(r"C:\写真\2026", "保存先を開く"),
            '=HYPERLINK("C:\\写真\\2026","保存先を開く")',
        )


class IntegratedCopyCsvTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source"
        self.destination = self.root / "保存先 folder"
        self.source.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def read_rows(self):
        with (self.destination / RESULT_LOG_NAME).open(
            "r", encoding="utf-8-sig", newline=""
        ) as handle:
            return list(csv.DictReader(handle))

    def make_plan(self, name="写真.txt"):
        original = self.source / name
        original.write_text("data", encoding="utf-8")
        plans = build_copy_plans(
            [detail(original, name)], self.destination, "{name}", False, "number"
        )
        return original, plans

    def test_copy_csv_and_shortcut_are_created(self):
        original, plans = self.make_plan()
        execute_copy_plans(plans, self.destination)

        rows = self.read_rows()
        self.assertEqual(rows[0]["処理結果"], "コピー完了")
        self.assertEqual(rows[0]["元ファイル名"], "写真.txt")
        self.assertTrue(rows[0]["保存先を開く"].startswith("=HYPERLINK("))
        self.assertEqual(rows[0]["元サイズ"], "4")
        self.assertEqual(rows[0]["保存後サイズ"], "4")
        shortcut = self.destination / "保存先を開く.lnk"
        fallback = self.destination / "保存先を開く.url"
        self.assertTrue(shortcut.exists() or fallback.exists())
        self.assertTrue((self.destination / original.name).exists())
        raw = (self.destination / RESULT_LOG_NAME).read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(raw.count(b"\xef\xbb\xbf"), 1)
        self.assertIn(b"\r\n", raw)

    def test_result_log_can_be_saved_outside_copy_destination(self):
        _, plans = self.make_plan("整理対象.txt")
        history_directory = self.root / "記録" / "履歴"
        execute_copy_plans(plans, self.destination, result_directory=history_directory)

        self.assertTrue((self.destination / "整理対象.txt").exists())
        self.assertTrue(
            (self.destination / "保存先を開く.lnk").exists()
            or (self.destination / "保存先を開く.url").exists()
        )
        self.assertTrue((history_directory / RESULT_LOG_NAME).exists())
        self.assertFalse((self.destination / RESULT_LOG_NAME).exists())

    def test_xlsx_option_appends_native_workbook_with_path_links(self):
        _, plans = self.make_plan("写真.xlsx")
        execute_copy_plans(plans, self.destination, output_format="xlsx")

        workbook_path = self.destination / "_DirectoryStructureGenerator_copy_log.xlsx"
        self.assertTrue(workbook_path.exists())
        workbook = load_workbook(workbook_path)
        sheet = workbook["コピー結果"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 3).value, "コピー完了")
        self.assertEqual(sheet.cell(2, 15).value, "保存先を開く")
        self.assertIn(str(self.destination), sheet.cell(2, 15).hyperlink.target)
        self.assertIn("写真.xlsx", sheet.cell(2, 16).hyperlink.target)

        _, second_plan = self.make_plan("二枚目.txt")
        execute_copy_plans(second_plan, self.destination, output_format="xlsx")
        workbook = load_workbook(workbook_path)
        self.assertEqual(workbook["コピー結果"].max_row, 3)

        _, reset_plan = self.make_plan("三枚目.txt")
        execute_copy_plans(reset_plan, self.destination, output_format="xlsx", history_mode="reset")
        workbook = load_workbook(workbook_path)
        self.assertEqual(workbook["コピー結果"].max_row, 2)
        archives = list(self.destination.glob("*_history_*.xlsx"))
        self.assertEqual(len(archives), 1)
        archived = load_workbook(archives[0])
        self.assertEqual(archived["コピー結果"].max_row, 3)

    def test_csv_reset_archives_prior_history_and_keeps_only_current_run_active(self):
        _, first_plan = self.make_plan("前回.txt")
        execute_copy_plans(first_plan, self.destination)
        _, second_plan = self.make_plan("今回.txt")
        execute_copy_plans(second_plan, self.destination, history_mode="reset")

        rows = self.read_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["元ファイル名"], "今回.txt")
        archives = list(self.destination.glob("*_history_*.csv"))
        self.assertEqual(len(archives), 1)
        with archives[0].open("r", encoding="utf-8-sig", newline="") as handle:
            archived_rows = list(csv.DictReader(handle))
        self.assertEqual(len(archived_rows), 1)
        self.assertEqual(archived_rows[0]["元ファイル名"], "前回.txt")

    @unittest.skipUnless(os.name != "nt", "PowerShell shortcut creation is covered by Windows builds")
    def test_windows_shortcut_uses_native_link_with_unicode_target(self):
        from dsg_app import organizer

        target = self.destination / "画像管理" / "写真"
        with patch.object(organizer, "os", SimpleNamespace(name="nt", environ=os.environ, fspath=os.fspath)), patch.object(
            organizer.subprocess, "run"
        ) as run:
            shortcut = organizer.create_destination_shortcut(target)

        self.assertEqual(shortcut.name, "保存先を開く.lnk")
        run.assert_called_once()
        self.assertEqual(run.call_args.kwargs["env"]["DSG_SHORTCUT_TARGET"], str(target.resolve()))

    @unittest.skipUnless(os.name != "nt", "PowerShell shortcut fallback is covered by Windows builds")
    def test_windows_shortcut_falls_back_to_url_when_com_is_unavailable(self):
        from dsg_app import organizer

        target = self.destination / "画像管理" / "写真"
        with patch.object(organizer, "os", SimpleNamespace(name="nt", environ=os.environ, fspath=os.fspath)), patch.object(
            organizer.subprocess, "run", side_effect=organizer.subprocess.CalledProcessError(1, "powershell")
        ):
            shortcut = organizer.create_destination_shortcut(target)

        self.assertEqual(shortcut.name, "保存先を開く.url")
        self.assertIn("file:///", shortcut.read_text(encoding="utf-8-sig"))

    def test_three_japanese_files_are_logged_and_originals_are_kept(self):
        names = ["写真 01.txt", "画像 二.txt", "資料 03.txt"]
        originals = []
        details = []
        for name in names:
            original = self.source / name
            original.write_text("data", encoding="utf-8")
            originals.append(original)
            details.append(detail(original, name))
        plans = build_copy_plans(details, self.destination, "{name}", False, "number")

        execute_copy_plans(plans, self.destination)

        rows = self.read_rows()
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(row["処理結果"] == "コピー完了" for row in rows))
        self.assertTrue(all(path.exists() for path in originals))
        self.assertTrue(all(plan.destination.exists() for plan in plans))

    def test_append_has_only_one_header(self):
        _, first = self.make_plan("first.txt")
        execute_copy_plans(first, self.destination)
        _, second = self.make_plan("second.txt")
        execute_copy_plans(second, self.destination)
        raw = (self.destination / RESULT_LOG_NAME).read_text(encoding="utf-8-sig")
        self.assertEqual(raw.count("実行ID,実行日時,処理結果"), 1)
        self.assertEqual(len(self.read_rows()), 2)

    def test_legacy_log_is_preserved_before_new_schema(self):
        self.destination.mkdir()
        (self.destination / RESULT_LOG_NAME).write_text(
            "日時,元ファイル,コピー先,結果,エラー\nold,row,,,\n", encoding="utf-8-sig"
        )
        _, plans = self.make_plan("new.txt")
        execute_copy_plans(plans, self.destination)
        legacy = list(self.destination.glob("*_legacy_*.csv"))
        self.assertEqual(len(legacy), 1)
        with (self.destination / RESULT_LOG_NAME).open(
            "r", encoding="utf-8-sig", newline=""
        ) as handle:
            self.assertEqual(next(csv.reader(handle)), RESULT_COLUMNS)

    def test_skip_result_is_written_without_overwriting_existing_file(self):
        original, _ = self.make_plan("collision.txt")
        self.destination.mkdir(parents=True)
        existing = self.destination / original.name
        existing.write_text("keep existing", encoding="utf-8")
        plans = build_copy_plans(
            [detail(original, original.name)], self.destination, "{name}", False, "skip"
        )

        execute_copy_plans(plans, self.destination)

        row = self.read_rows()[0]
        self.assertEqual(row["処理結果"], "既存のためスキップ")
        self.assertEqual(existing.read_text(encoding="utf-8"), "keep existing")
        self.assertTrue(original.exists())

    def test_error_row_is_written_and_source_is_kept(self):
        original, plans = self.make_plan("blocked.txt")
        with patch("dsg_app.organizer.shutil.copy2", side_effect=PermissionError("denied")):
            execute_copy_plans(plans, self.destination)
        row = self.read_rows()[0]
        self.assertEqual(row["処理結果"], "コピー失敗")
        self.assertEqual(row["エラー分類"], "ACCESS_DENIED")
        self.assertTrue(original.exists())

    def test_missing_source_is_logged_and_processing_continues(self):
        original, plans = self.make_plan("missing.txt")
        original.unlink()
        execute_copy_plans(plans, self.destination)
        self.assertEqual(self.read_rows()[0]["エラー分類"], "FILE_NOT_FOUND")

    def test_size_mismatch_never_deletes_source(self):
        original, plans = self.make_plan("important.txt")

        def bad_copy(_source, destination):
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_text("x", encoding="utf-8")

        with patch("dsg_app.organizer.shutil.copy2", side_effect=bad_copy):
            execute_copy_plans(plans, self.destination, delete_sources=True)
        self.assertTrue(original.exists())
        self.assertEqual(self.read_rows()[0]["処理結果"], "コピー完了・安全確認失敗")

    def test_formula_like_filename_is_neutralized(self):
        _, plans = self.make_plan("=SUM(A1).txt")
        execute_copy_plans(plans, self.destination)
        self.assertTrue(self.read_rows()[0]["元ファイル名"].startswith("'="))

    def test_all_formula_prefixes_are_neutralized(self):
        names = ["=formula.txt", "+formula.txt", "-formula.txt", "@formula.txt"]
        details = []
        for name in names:
            original = self.source / name
            original.write_text("data", encoding="utf-8")
            details.append(detail(original, name))
        plans = build_copy_plans(details, self.destination, "{name}", False, "number")

        execute_copy_plans(plans, self.destination)

        rows = self.read_rows()
        self.assertEqual(
            [row["元ファイル名"][0] for row in rows], ["'", "'", "'", "'"]
        )

    def test_cancel_is_logged(self):
        original, plans = self.make_plan("cancel.txt")
        with self.assertRaises(CopyCancelled):
            execute_copy_plans(plans, self.destination, cancel_requested=lambda: True)
        self.assertEqual(self.read_rows()[0]["エラー分類"], "CANCELLED")
        self.assertTrue(original.exists())

    def test_cancel_after_first_copy_records_unexecuted_rows(self):
        first, first_plans = self.make_plan("first.txt")
        second, second_plans = self.make_plan("second.txt")
        plans = [*first_plans, *second_plans]
        checks = iter([False, True])
        with self.assertRaises(CopyCancelled):
            execute_copy_plans(plans, self.destination, cancel_requested=lambda: next(checks))
        rows = self.read_rows()
        self.assertEqual([row["処理結果"] for row in rows], ["コピー完了", "中止により未実行"])
        self.assertTrue(first.exists())
        self.assertTrue(second.exists())
        self.assertTrue(plans[0].destination.exists())
        self.assertFalse(plans[1].destination.exists())

    def test_disk_full_copy_error_is_logged(self):
        original, plans = self.make_plan("disk-full.txt")
        error = OSError("disk full")
        error.errno = 28
        with patch("dsg_app.organizer.shutil.copy2", side_effect=error):
            execute_copy_plans(plans, self.destination)
        self.assertEqual(self.read_rows()[0]["エラー分類"], "DISK_FULL")
        self.assertTrue(original.exists())


class ErrorClassificationTests(unittest.TestCase):
    def test_disk_full_windows_error(self):
        exc = OSError("disk full")
        exc.winerror = 112
        self.assertEqual(classify_copy_error(exc), "DISK_FULL")


if __name__ == "__main__":
    unittest.main()
